#!/usr/bin/env python3
"""Construit les deux masters du kit presse à partir d'une charte validée.

Entrées :
  brand-profile.json  enrichi par brand_check.py (palettes + polices)
  plan-slides.<lang>.json et plan-document.<lang>.json (références de la skill par défaut)

Sorties (dans --out) :
  <Marque>_Presentation_presse_MASTER.pptx  16:9, thème, notes de consignes
  <Marque>_Communique_fiche_createurs_MASTER.docx  fond blanc, styles nommés
  champs.json  liste de tous les champs {{...}} présents dans les deux masters

Les champs {{MARQUE}} et {{SITE}} sont remplis depuis la charte ; tous les autres
restent à remplir à chaque copie du master.

Usage :
  python3 build_masters.py brand-profile.json --out masters/ [--logo logo.png]
         [--logo-clair logo-pour-fond-clair.png] [--slides-plan p.json] [--doc-plan p.json]
         [--only slides|document]
"""
import argparse
import copy
import json
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt as DPt, RGBColor as DRGB

HERE = os.path.dirname(os.path.abspath(__file__))
REFS = os.path.join(os.path.dirname(HERE), "references")
FIELD = re.compile(r"\{\{[A-Z0-9_]+\}\}")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def hexrgb(h, cls=RGBColor):
    h = h.lstrip("#")
    return cls(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def slug(text):
    text = re.sub(r"[^A-Za-z0-9]+", "_", text or "Marque").strip("_")
    return text or "Marque"


class Ctx:
    def __init__(self, profile, args):
        self.p = profile
        pal = profile.get("palettes")
        if not pal:
            sys.exit("Palettes absentes : lancer d'abord brand_check.py sur ce profil.")
        self.s = pal["slides"]
        self.d = pal["document"]
        fonts = profile.get("polices") or {}
        self.f_title = (fonts.get("titres") or {}).get("famille") or "Arial"
        self.f_text = (fonts.get("texte") or {}).get("famille") or self.f_title
        self.fills = {"{{MARQUE}}": profile.get("marque") or "", "{{SITE}}": profile.get("site") or ""}
        self.fills = {k: v for k, v in self.fills.items() if v}
        self.logo = args.logo
        self.logo_light = args.logo_clair or args.logo
        self.fields = set()

    def t(self, text):
        for k, v in self.fills.items():
            text = text.replace(k, v)
        self.fields.update(FIELD.findall(text))
        return text


# ---------------------------------------------------------------- slides

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.8)


def set_theme(prs, c):
    master = prs.slide_master
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": A}
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", "Marque")
    values = {
        "dk1": c.s["fond"] if c.p["palettes"]["theme_slides"] == "clair" else c.s["texte"],
        "lt1": c.s["fond"] if c.p["palettes"]["theme_slides"] == "sombre" else c.s["texte"],
        "dk2": c.s["surface"], "lt2": c.s["texte_secondaire"],
        "accent1": c.s["primaire"], "accent2": c.s["secondaire"], "accent3": c.s["accent"],
        "accent4": c.s["bordure"], "accent5": c.d["titres"], "accent6": c.d["filet"],
        "hlink": c.s["accent"], "folHlink": c.s["secondaire"],
    }
    # Fond sombre : dk1/lt1 inversés pour que le texte par défaut reste lisible.
    if c.p["palettes"]["theme_slides"] == "sombre":
        values["dk1"], values["lt1"] = c.s["texte"], c.s["fond"]
    for tag, value in values.items():
        el = scheme.find("a:" + tag, ns)
        if el is None:
            continue
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, "{%s}srgbClr" % A).set("val", value.lstrip("#"))
    fs = root.find(".//a:fontScheme", ns)
    fs.set("name", "Marque")
    for tag, face in (("majorFont", c.f_title), ("minorFont", c.f_text)):
        latin = fs.find("a:%s/a:latin" % tag, ns)
        if latin is not None:
            latin.set("typeface", face)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    fill = master.background.fill
    fill.solid()
    fill.fore_color.rgb = hexrgb(c.s["fond"])
    for layout in master.slide_layouts:
        lf = layout.background.fill
        lf.solid()
        lf.fore_color.rgb = hexrgb(c.s["fond"])


def text_box(slide, c, x, y, w, h, text, size, color, font=None, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, name=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = c.t(line)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font or c.f_text
        run.font.color.rgb = hexrgb(color)
    return box


def bullets(slide, c, x, y, w, h, items, size=22):
    box = slide.shapes.add_textbox(x, y, w, h)
    box.name = "Puces"
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(14)
        mark = para.add_run()
        mark.text = "■  "
        mark.font.size = Pt(size * 0.6)
        mark.font.color.rgb = hexrgb(c.s["accent"])
        mark.font.name = c.f_text
        run = para.add_run()
        run.text = c.t(item)
        run.font.size = Pt(size)
        run.font.name = c.f_text
        run.font.color.rgb = hexrgb(c.s["texte"])


def rect(slide, c, x, y, w, h, fill=None, line=None, dashed=False, shape=MSO_SHAPE.RECTANGLE, name=None):
    r = slide.shapes.add_shape(shape, x, y, w, h)
    if name:
        r.name = name
    if fill:
        r.fill.solid()
        r.fill.fore_color.rgb = hexrgb(fill)
    else:
        r.fill.background()
    if line:
        r.line.color.rgb = hexrgb(line)
        r.line.width = Pt(1.5)
        if dashed:
            r.line.dash_style = MSO_LINE.DASH
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def image_frame(slide, c, x, y, w, h, label):
    r = rect(slide, c, x, y, w, h, fill=c.s["surface"], line=c.s["accent"], dashed=True, name="Image : " + label)
    tf = r.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = "[Visuel] " + label
    run.font.size = Pt(14)
    run.font.name = c.f_text
    run.font.color.rgb = hexrgb(c.s["texte_secondaire"])
    return r


def logo(slide, c, x, y, h, light=False):
    path = c.logo_light if light else c.logo
    if path and os.path.exists(path):
        pic = slide.shapes.add_picture(path, x, y, height=h)
        pic.name = "Logo"
        return pic
    return image_frame(slide, c, x, y, Inches(2.2), h, "Logo")


def header(slide, c, kicker, title):
    rect(slide, c, MX, Inches(0.62), Inches(0.6), Inches(0.08), fill=c.s["accent"], name="Filet")
    if kicker:
        text_box(slide, c, MX, Inches(0.78), Inches(10), Inches(0.45), kicker.upper(), 14,
                 c.s["primaire"], font=c.f_title, bold=True, name="Surtitre")
    if title:
        text_box(slide, c, MX, Inches(1.2), W - 2 * MX, Inches(1.2), title, 36, c.s["texte"],
                 font=c.f_title, bold=True, name="Titre")


def footer(slide, c, plan, number):
    text_box(slide, c, MX, H - Inches(0.55), Inches(8), Inches(0.3), plan.get("pied", ""), 10,
             c.s["texte_secondaire"], name="Pied")
    text_box(slide, c, W - MX - Inches(1), H - Inches(0.55), Inches(1), Inches(0.3), str(number), 10,
             c.s["texte_secondaire"], align=PP_ALIGN.RIGHT, name="Numéro")


def render_slide(prs, c, plan, spec, number, all_fields):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    lay = spec["layout"]
    top = Inches(2.6)
    body_h = H - top - Inches(0.9)

    if lay == "guide":
        rect(slide, c, 0, 0, W, Inches(0.25), fill=c.s["accent"])
        header(slide, c, "Réservé à l'équipe", spec["titre"])
        bullets(slide, c, MX, top - Inches(0.2), Inches(6.6), body_h, spec["puces"], size=18)
        rect(slide, c, Inches(7.9), top - Inches(0.2), W - Inches(7.9) - MX, Inches(2.6), fill=c.s["surface"], line=c.s["bordure"])
        text_box(slide, c, Inches(8.2), top, Inches(4.2), Inches(0.4), "CHAMPS DU MASTER", 12,
                 c.s["primaire"], font=c.f_title, bold=True)
        text_box(slide, c, Inches(8.2), top + Inches(0.5), Inches(4.2), Inches(1.7),
                 ["%d champs {{...}} à remplir." % len(all_fields),
                  "Liste complète : notes de cette slide et fichier champs.json livré avec le master."],
                 14, c.s["texte_secondaire"])
    elif lay == "cover":
        rect(slide, c, 0, H - Inches(0.18), W, Inches(0.18), fill=c.s["accent"])
        logo(slide, c, MX, Inches(0.7), Inches(0.9))
        text_box(slide, c, MX, Inches(2.5), Inches(11), Inches(0.5), spec.get("surtitre", "").upper(), 16,
                 c.s["primaire"], font=c.f_title, bold=True, name="Surtitre")
        text_box(slide, c, MX, Inches(3.0), Inches(11.5), Inches(1.6), spec["titre"], 66, c.s["texte"],
                 font=c.f_title, bold=True, name="Titre")
        text_box(slide, c, MX, Inches(4.7), Inches(10.5), Inches(1.2), spec.get("sous_titre", ""), 26,
                 c.s["texte_secondaire"], name="Promesse")
        text_box(slide, c, MX, H - Inches(0.85), Inches(8), Inches(0.4), plan.get("pied", ""), 12, c.s["texte_secondaire"])
    elif lay == "statement":
        header(slide, c, spec.get("surtitre"), None)
        text_box(slide, c, MX, Inches(1.9), W - 2 * MX, Inches(4.2), spec["texte"], 44, c.s["texte"],
                 font=c.f_title, bold=True, anchor=MSO_ANCHOR.MIDDLE, name="Phrase")
    elif lay == "bullets":
        header(slide, c, spec.get("surtitre"), spec.get("titre"))
        if spec.get("image"):
            bullets(slide, c, MX, top, Inches(6.2), body_h, spec["puces"])
            image_frame(slide, c, Inches(7.4), top, W - Inches(7.4) - MX, body_h, spec["image"])
        else:
            bullets(slide, c, MX, top, W - 2 * MX, body_h, spec["puces"], size=24)
    elif lay in ("cards", "steps", "numbers", "gallery"):
        header(slide, c, spec.get("surtitre"), spec.get("titre"))
        items = spec.get("cartes") or spec.get("etapes") or spec.get("chiffres") or spec.get("images")
        n = len(items)
        gap = Inches(0.35)
        cw = int((W - 2 * MX - gap * (n - 1)) / n)
        for i, item in enumerate(items):
            x = MX + (cw + gap) * i
            if lay == "cards":
                rect(slide, c, x, top, cw, body_h, fill=c.s["surface"], line=c.s["bordure"], name="Carte")
                rect(slide, c, x, top, cw, Inches(0.08), fill=c.s["accent"])
                text_box(slide, c, x + Inches(0.3), top + Inches(0.35), cw - Inches(0.6), Inches(1.1), item["titre"], 20,
                         c.s["texte"], font=c.f_title, bold=True)
                text_box(slide, c, x + Inches(0.3), top + Inches(1.5), cw - Inches(0.6), body_h - Inches(1.8), item["texte"], 16,
                         c.s["texte_secondaire"])
            elif lay == "steps":
                text_box(slide, c, x, top, cw, Inches(1.3), "%02d" % (i + 1), 60, c.s["primaire"], font=c.f_title, bold=True)
                rect(slide, c, x, top + Inches(1.45), cw, Inches(0.04), fill=c.s["bordure"])
                text_box(slide, c, x, top + Inches(1.7), cw, body_h - Inches(1.7), item, 20, c.s["texte"])
            elif lay == "numbers":
                text_box(slide, c, x, top + Inches(0.2), cw, Inches(1.6), item["valeur"], 44, c.s["primaire"],
                         font=c.f_title, bold=True)
                text_box(slide, c, x, top + Inches(1.9), cw, body_h - Inches(1.9), item["legende"], 18, c.s["texte_secondaire"])
            else:
                image_frame(slide, c, x, top, cw, body_h - Inches(0.8), item["cadre"])
                text_box(slide, c, x, top + body_h - Inches(0.65), cw, Inches(0.6), item["legende"], 14, c.s["texte_secondaire"])
    elif lay == "quote":
        header(slide, c, spec.get("surtitre"), None)
        text_box(slide, c, MX, Inches(1.5), Inches(1.5), Inches(1.5), "“", 120, c.s["accent"], font=c.f_title, bold=True)
        text_box(slide, c, MX + Inches(0.2), Inches(2.6), Inches(8), Inches(2.6), spec["citation"], 30, c.s["texte"],
                 font=c.f_title, name="Citation")
        text_box(slide, c, MX + Inches(0.2), Inches(5.4), Inches(8), Inches(0.6), spec["auteur"], 18, c.s["primaire"],
                 bold=True, name="Auteur")
        image_frame(slide, c, Inches(9.6), Inches(1.9), Inches(2.9), Inches(3.9), spec.get("image", "Photo"))
    elif lay == "contact":
        rect(slide, c, 0, H - Inches(0.18), W, Inches(0.18), fill=c.s["accent"])
        text_box(slide, c, MX, Inches(1.6), Inches(8), Inches(1.2), spec["titre"], 48, c.s["texte"], font=c.f_title, bold=True)
        text_box(slide, c, MX, Inches(3.0), Inches(8), Inches(3), spec["lignes"], 22, c.s["texte_secondaire"])
        logo(slide, c, Inches(9.5), Inches(3.0), Inches(1.2))
    else:
        sys.exit("Layout inconnu : %s" % lay)

    if lay not in ("cover", "contact", "guide"):
        footer(slide, c, plan, number)
    notes = spec.get("notes")
    if lay == "guide":
        notes = (notes or "") + "\n\nChamps du master :\n" + "\n".join(sorted(all_fields))
    if notes:
        slide.notes_slide.notes_text_frame.text = c.t(notes)
    return slide


def build_slides(c, plan, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    set_theme(prs, c)
    # Premier passage pour connaître tous les champs (affichés sur la slide « mode d'emploi »).
    probe = Ctx.__new__(Ctx)
    probe.__dict__.update(c.__dict__)
    probe.fields = set()
    probe.t(json.dumps(plan, ensure_ascii=False))
    all_fields = probe.fields
    for i, spec in enumerate(plan["slides"]):
        render_slide(prs, c, plan, spec, i, all_fields)
    prs.core_properties.title = c.t(plan.get("titre_fichier", "Présentation presse - MASTER"))
    path = os.path.join(out, "%s_Presentation_presse_MASTER.pptx" % slug(c.p.get("marque")))
    prs.save(path)
    return path


# ---------------------------------------------------------------- document

def shade(cell, color):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color.lstrip("#"))
    tc.append(sh)


def border_left(paragraph, color, size=24):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "12")
    left.set(qn("w:color"), color.lstrip("#"))
    bdr.append(left)
    ppr.append(bdr)


def style(doc, name, base, font, size, color, bold=False, italic=False, before=0, after=6, upper=False):
    styles = doc.styles
    st = styles[name] if name in [s.name for s in styles] else styles.add_style(name, 1)
    if base:
        st.base_style = styles[base]
    st.font.name = font
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = DPt(size)
    st.font.color.rgb = hexrgb(color, DRGB)
    st.font.bold = bold
    st.font.italic = italic
    st.font.all_caps = upper
    st.paragraph_format.space_before = DPt(before)
    st.paragraph_format.space_after = DPt(after)
    return st


def build_document(c, plan, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    d = c.d
    style(doc, "Normal", None, c.f_text, 11, d["texte"], after=6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    style(doc, "Title", "Normal", c.f_title, 24, d["titres"], bold=True, before=4, after=8)
    style(doc, "Heading 2", "Normal", c.f_title, 14, d["titres"], bold=True, before=12, after=4)
    style(doc, "Heading 3", "Normal", c.f_title, 12, d["texte"], bold=True, before=8, after=3)
    style(doc, "Chapeau", "Normal", c.f_text, 12.5, d["texte"], bold=True, after=10)
    style(doc, "Etiquette", "Normal", c.f_title, 10, d["filet"] if d["filet"] != d["fond"] else d["titres"], bold=True, upper=True, after=2)
    style(doc, "Citation marque", "Normal", c.f_title, 13, d["texte"], italic=True, before=10, after=2)
    style(doc, "Auteur citation", "Normal", c.f_text, 10.5, d["titres"], bold=True, after=10)
    style(doc, "Consigne", "Normal", c.f_text, 9, d["texte_secondaire"], italic=True, after=8)
    style(doc, "List Bullet", "Normal", c.f_text, 11, d["texte"], after=3)

    footer_p = sec.footer.paragraphs[0]
    footer_p.text = c.t("{{MARQUE}} · {{SITE}}")
    footer_p.style = doc.styles["Consigne"]

    body = doc.element.body
    # Retire le paragraphe vide initial.
    for p in list(doc.paragraphs):
        if not p.text:
            body.remove(p._p)

    def para(text, st="Normal"):
        return doc.add_paragraph(c.t(text), style=st)

    for b in plan["blocs"]:
        t = b["type"]
        if t == "logo":
            p = doc.add_paragraph()
            if c.logo_light and os.path.exists(c.logo_light):
                p.add_run().add_picture(c.logo_light, height=Cm(1.4))
            else:
                r = p.add_run("[Logo de la marque, version pour fond clair]")
                r.font.color.rgb = hexrgb(d["texte_secondaire"], DRGB)
                r.font.italic = True
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("val", "single"), ("sz", "12"), ("space", "8"), ("color", d["filet"].lstrip("#"))):
                bottom.set(qn("w:" + k), v)
            border.append(bottom)
            p._p.get_or_add_pPr().append(border)
            p.paragraph_format.space_after = DPt(12)
        elif t == "etiquette":
            para(b["texte"], "Etiquette")
        elif t == "titre":
            para(b["texte"], "Title")
        elif t == "chapeau":
            para(b["texte"], "Chapeau")
        elif t in ("h2", "h3"):
            para(b["texte"], "Heading 2" if t == "h2" else "Heading 3")
        elif t == "para":
            para(b["texte"])
        elif t == "consigne":
            para("Consigne : " + b["texte"], "Consigne")
        elif t == "puces":
            for item in b["items"]:
                para(item, "List Bullet")
        elif t == "citation":
            p = para("« " + b["texte"] + " »", "Citation marque")
            border_left(p, d["filet"])
            p2 = para(b["auteur"], "Auteur citation")
            border_left(p2, d["filet"])
        elif t == "encadre":
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.rows[0].cells[0]
            shade(cell, d["encadre_fond"])
            cell.paragraphs[0].text = c.t(b["titre"])
            cell.paragraphs[0].style = doc.styles["Heading 3"]
            for item in b["items"]:
                cell.add_paragraph(c.t(item), style="List Bullet")
            doc.add_paragraph()
        elif t == "tableau":
            table = doc.add_table(rows=1 + len(b["lignes"]), cols=len(b["entetes"]))
            table.style = doc.styles["Table Grid"]
            for j, h in enumerate(b["entetes"]):
                cell = table.rows[0].cells[j]
                shade(cell, d["encadre_fond"])
                cell.paragraphs[0].text = c.t(h)
                cell.paragraphs[0].runs[0].font.bold = True
            for i, row in enumerate(b["lignes"]):
                for j, v in enumerate(row):
                    table.rows[i + 1].cells[j].paragraphs[0].text = c.t(v)
            doc.add_paragraph()
        elif t == "qr":
            para(b["question"], "Heading 3")
            para(b["reponse"])
        elif t == "saut":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            sys.exit("Bloc inconnu : %s" % t)

    doc.core_properties.title = c.t(plan.get("titre_fichier", "Communiqué - MASTER"))
    path = os.path.join(out, "%s_Communique_fiche_createurs_MASTER.docx" % slug(c.p.get("marque")))
    doc.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("profile")
    ap.add_argument("--out", default="masters")
    ap.add_argument("--logo", help="Logo pour le fond des slides")
    ap.add_argument("--logo-clair", help="Logo lisible sur fond blanc (document)")
    ap.add_argument("--slides-plan", default=os.path.join(REFS, "plan-slides.fr.json"))
    ap.add_argument("--doc-plan", default=os.path.join(REFS, "plan-document.fr.json"))
    ap.add_argument("--only", choices=("slides", "document"))
    args = ap.parse_args()

    with open(args.profile, encoding="utf-8") as f:
        profile = json.load(f)
    os.makedirs(args.out, exist_ok=True)
    c = Ctx(profile, args)
    made = []
    if args.only != "document":
        with open(args.slides_plan, encoding="utf-8") as f:
            made.append(build_slides(c, json.load(f), args.out))
    if args.only != "slides":
        with open(args.doc_plan, encoding="utf-8") as f:
            made.append(build_document(c, json.load(f), args.out))
    fields_path = os.path.join(args.out, "champs.json")
    with open(fields_path, "w", encoding="utf-8") as f:
        json.dump({"remplis_depuis_la_charte": c.fills, "a_remplir": sorted(c.fields)}, f, ensure_ascii=False, indent=2)
    for m in made + [fields_path]:
        print(m)
    print("%d champs à remplir." % len(c.fields))


if __name__ == "__main__":
    main()
